"""Convert monchmonch_model_v6.xlsx INPUTS tab → calc scenario JSON.

Usage:
    python3 scripts/xlsx_to_scenario.py <xlsx> [out.json]

Then in the live calc (monchmonch-calc.vercel.app):
    1. Click "Load Scenario"
    2. Pick the generated .json file
    3. All inputs update; P&L + cash flow recompute

Cell map mirrors AUDIT_SYNC.md.
"""

import json
import sys
from pathlib import Path
from openpyxl import load_workbook


def read_cell(ws, addr, default=None):
    v = ws[addr].value
    return v if v is not None else default


def read_range_horizontal(ws, row, start_col, count):
    """Read N cells in a row, returning a list."""
    return [ws.cell(row=row, column=start_col + i).value for i in range(count)]


def xlsx_to_scenario(xlsx_path: str) -> dict:
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb["INPUTS"]

    state = {}

    # === GLOBAL ASSUMPTIONS (R5-R14) ===
    state["returnsPct"] = read_cell(ws, "B5")
    state["targetDOI"] = read_cell(ws, "B6")
    state["safetyStockPct"] = read_cell(ws, "B7")
    state["warehousingPerUnit"] = read_cell(ws, "B8")
    state["insurancePct"] = read_cell(ws, "B9")
    state["costOfCapitalPct"] = read_cell(ws, "B10")
    state["spoilageBar"] = read_cell(ws, "B11")
    state["spoilageElec"] = read_cell(ws, "B12")
    state["spoilageCostMult"] = read_cell(ws, "B13")
    state["carryingCostRate"] = read_cell(ws, "B14")

    # === LABOR & OVERHEAD (R18-R27) ===
    state["laborRoles"] = []
    for r in range(18, 22):
        role = read_cell(ws, f"A{r}")
        if role:
            state["laborRoles"].append({
                "role": role,
                "headcount": read_cell(ws, f"B{r}"),
                "rate": read_cell(ws, f"C{r}"),
                "hoursPerRun": read_cell(ws, f"D{r}"),
                "isVariable": read_cell(ws, f"E{r}") == "Yes",
            })
    state["laborPerUnit"] = read_cell(ws, "B24")
    state["fixedOverhead"] = read_cell(ws, "B25")
    state["equipAmort"] = read_cell(ws, "B26")
    state["gna"] = read_cell(ws, "B27")

    # === LEGACY PRICING (R30-R37) ===
    state["barMSRP"] = read_cell(ws, "B30")
    state["barTrade"] = read_cell(ws, "B31")
    state["elecMSRP"] = read_cell(ws, "B32")
    state["elecTrade"] = read_cell(ws, "B33")
    state["fulfillCost"] = read_cell(ws, "B34")
    state["freightCost"] = read_cell(ws, "B35")
    state["marketingPct"] = read_cell(ws, "B36")
    state["marketingAnnual"] = read_cell(ws, "B37")

    # === CHANNELS (R41-R46) ===
    state["channels"] = []
    for r in range(41, 47):
        name = read_cell(ws, f"A{r}")
        if name:
            state["channels"].append({
                "name": name,
                "active": read_cell(ws, f"B{r}") == "Yes",
                "pctAlloc": read_cell(ws, f"C{r}") or 0,
                "barPrice": read_cell(ws, f"D{r}"),
                "elecPrice": read_cell(ws, f"E{r}"),
                "costPerUnit": read_cell(ws, f"F{r}"),
                "rampMonths": int(read_cell(ws, f"G{r}") or 0),
                "maxMonthlyUnits": int(read_cell(ws, f"H{r}") or 0),
            })

    # === CO-MAN PRICING (R51-R57) ===
    state["coManThresholds"] = []
    state["coManLabels"] = []
    state["coManBars"] = []
    state["coManElec"] = []
    labor_scales = []
    for r in range(51, 58):
        state["coManLabels"].append(read_cell(ws, f"B{r}"))
        state["coManThresholds"].append(read_cell(ws, f"C{r}") or 0)
        state["coManBars"].append(read_cell(ws, f"E{r}"))
        state["coManElec"].append(read_cell(ws, f"F{r}"))
        labor_scales.append(read_cell(ws, f"G{r}"))

    # === GROWTH + ANNUAL OPEX ARRAYS (R61-R67) ===
    growth_y = read_range_horizontal(ws, 61, 2, 5)  # B-F
    state["growthY2"] = growth_y[1]
    state["growthY3"] = growth_y[2]
    state["growthY4"] = growth_y[3]
    state["growthY5"] = growth_y[4]
    state["marketingByYear"] = read_range_horizontal(ws, 62, 2, 5)
    state["payrollByYear"] = read_range_horizontal(ws, 63, 2, 5)
    state["bdSalesByYear"] = read_range_horizontal(ws, 64, 2, 5)
    state["clinicalByYear"] = read_range_horizontal(ws, 65, 2, 5)
    state["overheadMult"] = read_range_horizontal(ws, 66, 2, 5)
    state["licensingByYear"] = read_range_horizontal(ws, 67, 2, 5)
    state["licensingCogsPct"] = read_cell(ws, "B68")
    state["payrollY1"] = read_cell(ws, "B69")
    state["payrollStartMonth"] = int(read_cell(ws, "B70") or 1)

    # === BAR RM (R74-R87) ===
    state["barRM"] = []
    for r in range(74, 88):
        name = read_cell(ws, f"A{r}")
        if name:
            state["barRM"].append({
                "name": name,
                "qty": read_cell(ws, f"B{r}") or 0,
                "t1": read_cell(ws, f"C{r}") or 0,
                "t2": read_cell(ws, f"D{r}") or 0,
                "t3": read_cell(ws, f"E{r}") or 0,
                "moq": int(read_cell(ws, f"F{r}") or 0),
            })

    # === ELEC RM (R91-R99) ===
    state["elecRM"] = []
    for r in range(91, 100):
        name = read_cell(ws, f"A{r}")
        if name:
            state["elecRM"].append({
                "name": name,
                "qty": read_cell(ws, f"B{r}") or 0,
                "t1": read_cell(ws, f"C{r}") or 0,
                "t2": read_cell(ws, f"D{r}") or 0,
                "t3": read_cell(ws, f"E{r}") or 0,
                "moq": int(read_cell(ws, f"F{r}") or 0),
            })

    # === Y1 DEMAND + SEASONALITY (R103-R114) ===
    state["barDemand"] = [int(read_cell(ws, f"B{r}") or 0) for r in range(103, 115)]
    state["elecDemand"] = [int(read_cell(ws, f"C{r}") or 0) for r in range(103, 115)]
    state["barSeason"] = [read_cell(ws, f"D{r}") or 1.0 for r in range(103, 115)]
    state["elecSeason"] = [read_cell(ws, f"E{r}") or 1.0 for r in range(103, 115)]

    # === LAUNCH EXPENSES (R119-R124) ===
    state["launchExpenses"] = []
    for r in range(119, 125):
        name = read_cell(ws, f"A{r}")
        if name:
            state["launchExpenses"].append({
                "name": name,
                "cost": read_cell(ws, f"B{r}") or 0,
                "month": int(read_cell(ws, f"C{r}") or 0),
                "notes": read_cell(ws, f"D{r}") or "",
            })

    # === FUNDING (R128-R132) ===
    state["startingCash"] = read_cell(ws, "B128")
    state["equityRaised"] = read_cell(ws, "B129")
    state["debtAmount"] = read_cell(ws, "B130")
    state["debtRate"] = read_cell(ws, "B131")
    state["debtTerm"] = int(read_cell(ws, "B132") or 0)

    # === COMMITMENTS (R136-R138) ===
    state["commitments"] = []
    for r in range(136, 139):
        name = read_cell(ws, f"A{r}")
        if name:
            state["commitments"].append({
                "name": name,
                "type": read_cell(ws, f"B{r}") or "",
                "monthlyCost": read_cell(ws, f"C{r}") or 0,
                "termMonths": int(read_cell(ws, f"D{r}") or 0),
                "notes": read_cell(ws, f"E{r}") or "",
            })

    # === MIN RUNS (R142-R143) ===
    state["minBarRun"] = int(read_cell(ws, "B142") or 0)
    state["minElecRun"] = int(read_cell(ws, "B143") or 0)

    # === WORKING CAPITAL (R147-R149, added in v6) ===
    state["arDays"] = read_cell(ws, "B147") or 35
    state["inventoryDays"] = read_cell(ws, "B148") or 60
    state["apDays"] = read_cell(ws, "B149") or 30

    # === Defaults for fields the spreadsheet doesn't carry but calc expects ===
    state.setdefault("payrollAnnualIncrease", 200000)
    state.setdefault("rmLeadWeeks", 4)
    state.setdefault("coManLeadWeeks", 3)
    state.setdefault("rmPaymentTerms", 30)
    state.setdefault("coManPaymentTerms", 30)
    state.setdefault("coManDeposit", 0.25)

    # Packaging arrays — spreadsheet doesn't have these as INPUTS rows
    state.setdefault("barPackaging", [
        {"name": "Wrapper/Film", "costPerUnit": 0, "moq": 50000, "leadWeeks": 6},
        {"name": "Display Box (12ct)", "costPerUnit": 0, "moq": 5000, "leadWeeks": 4},
        {"name": "Labels", "costPerUnit": 0, "moq": 25000, "leadWeeks": 3},
        {"name": "Inserts", "costPerUnit": 0, "moq": 25000, "leadWeeks": 3},
    ])
    state.setdefault("elecPackaging", [
        {"name": "Foil Sachet", "costPerUnit": 0, "moq": 100000, "leadWeeks": 8},
        {"name": "Display Box (20ct)", "costPerUnit": 0, "moq": 5000, "leadWeeks": 4},
        {"name": "Labels", "costPerUnit": 0, "moq": 25000, "leadWeeks": 3},
    ])
    state.setdefault("fundingSources", [
        {"name": "Founder Equity (sweat + cash)", "amount": 100000, "type": "equity"},
        {"name": "Seed Round (current ask)", "amount": 2500000, "type": "equity"},
    ])

    return state


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)

    xlsx_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2]) if len(sys.argv) > 2 else xlsx_path.with_suffix(".scenario.json")

    if not xlsx_path.exists():
        print(f"ERROR: {xlsx_path} not found", file=sys.stderr)
        sys.exit(1)

    state = xlsx_to_scenario(str(xlsx_path))
    out_path.write_text(json.dumps(state, indent=2))

    bar_rm_count = len([m for m in state["barRM"] if m["qty"] > 0])
    print(f"✓ Wrote {out_path}")
    print(f"  - {len(state['channels'])} channels ({sum(1 for c in state['channels'] if c['active'])} active)")
    print(f"  - {bar_rm_count}/{len(state['barRM'])} bar RM ingredients populated")
    print(f"  - {len(state['elecRM'])} elec RM ingredients")
    print(f"  - Seed: ${state['equityRaised']:,.0f} | WC: AR {state['arDays']}d / Inv {state['inventoryDays']}d / AP {state['apDays']}d")
    print(f"  - Marketing Y1-Y5: {[f'${int(m):,}' for m in state['marketingByYear']]}")
    print(f"  - Licensing Y1-Y5: {[f'${int(m):,}' for m in state['licensingByYear']]}")
    print(f"\nLoad into calc:")
    print(f"  1. Open https://monchmonch-calc.vercel.app/")
    print(f"  2. Click 'Load Scenario'")
    print(f"  3. Pick {out_path}")


if __name__ == "__main__":
    main()
